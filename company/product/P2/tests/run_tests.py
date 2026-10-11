#!/usr/bin/env python3
"""原価率計算シートの計算テスト。

1. tests/cases.py の各ケースの入力で、build.py と同じ数式のブックを作る。
2. LibreOffice（soffice --headless --convert-to xlsx）で開いて保存し直し、数式を再計算させる。
3. 再計算後のセル値を読み、(a) ケースに手で書いた期待値、(b) model.py で独立に計算した全セルの期待値と照合する。
4. 結果を tests/result.md に書く。不合格があれば終了コード 1。

使い方:  python3 company/product/P2/tests/run_tests.py
"""
from __future__ import annotations

import datetime
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
P2 = os.path.dirname(HERE)
sys.path.insert(0, P2)
sys.path.insert(0, HERE)

import build  # noqa: E402
import model  # noqa: E402
from cases import CASES, S, sample_case  # noqa: E402

ERROR_STRINGS = ("#DIV/0!", "#N/A", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "Err:")


def recalc(paths, outdir):
    cmd = ["soffice", "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", outdir] + paths
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=600)


def norm(v):
    return "" if v is None else v


def same(actual, expected):
    a, e = norm(actual), norm(expected)
    if isinstance(e, (int, float)) and not isinstance(e, bool):
        return isinstance(a, (int, float)) and abs(a - e) < 1e-9
    return str(a) == str(e)


def addr(key):
    """check キー → (シート名, セル番地)"""
    p = key.split(":")
    if p[0] == "mat":
        return build.MAT_SHEET, f"{p[2]}{build.MAT_FIRST + int(p[1]) - 1}"
    if p[0] == "menu":
        c = build.menu_cells(int(p[1]) - 1)
        if p[2] == "line":
            return build.MENU_SHEET, f"{p[4]}{c['line_first'] + int(p[3]) - 1}"
        return build.MENU_SHEET, c[p[2]]
    if p[0] == "list":
        col = {k: col for col, _, _, _, k in build.LIST_COLS[1:]}[p[2]]
        return build.LIST_SHEET, f"{col}{build.LIST_FIRST + int(p[1]) - 1}"
    raise ValueError(key)


def model_checks(case):
    """model.py の独立計算から、全ての計算セルの期待値を作る"""
    settings = case.get("settings", S)
    mat_rows, menus = model.compute(case["materials"], case["menus"], settings)
    out = []
    for i, row in enumerate(mat_rows, start=1):
        for col in "JKLM":
            out.append((f"mat:{i}:{col}", row[col]))
    for i, m in enumerate(menus, start=1):
        for k, ln in enumerate(m["lines"], start=1):
            for col in "DEFG":
                out.append((f"menu:{i}:line:{k}:{col}", ln[col]))
        for key in ("total", "target", "tax", "net", "rate", "diff", "judge", "rec"):
            out.append((f"menu:{i}:{key}", m[key]))
        for key in ("price", "total", "rate", "target", "diff", "judge", "rec"):
            v = m[key] if not (key == "price" and m["price"] is None) else ""
            out.append((f"list:{i}:{key}", v if m["name"] else ""))
        out.append((f"list:{i}:name", m["name"] or ""))
    return out


def scan_errors(wb):
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERROR_STRINGS):
                    errs.append(f"{ws.title}!{c.coordinate}={c.value}")
    return errs


def main():
    cases = list(CASES) + [sample_case()]
    work = tempfile.mkdtemp(prefix="p2test_")
    src, out = os.path.join(work, "src"), os.path.join(work, "out")
    os.makedirs(src)
    os.makedirs(out)
    paths = []
    for case in cases:
        p = os.path.join(src, f"{case['id']}.xlsx")
        if case.get("use_dist"):
            shutil.copy(build.OUT, p)   # 同梱ファイルそのものを再計算する
        else:
            build.build_workbook(p, case["materials"], case["menus"], case.get("settings", S))
        paths.append(p)
    recalc(paths, out)
    ver = subprocess.run(["soffice", "--version"], stdout=subprocess.PIPE, text=True).stdout.strip()

    results = []
    total_pass = total_fail = 0
    for case in cases:
        wb = openpyxl.load_workbook(os.path.join(out, f"{case['id']}.xlsx"), data_only=True)
        hand, auto = case["checks"], model_checks(case)
        fails = []
        n = 0
        # 手計算の期待値と model の期待値が食い違っていないか（model の検算）
        auto_map = dict(auto)
        for key, exp in hand:
            if key in auto_map and not same(auto_map[key], exp):
                fails.append(f"model不一致 {key}: 手計算={exp!r} model={auto_map[key]!r}")
        for key, exp in hand + auto:
            sh, ref = addr(key)
            act = wb[sh][ref].value
            n += 1
            if not same(act, exp):
                fails.append(f"{key}（{sh}!{ref}）: 期待={exp!r} 実際={act!r}")
        errs = scan_errors(wb)
        fails += [f"エラー値: {e}" for e in errs]
        ok = not fails
        total_pass += ok
        total_fail += (not ok)
        sample_vals = None
        if case.get("use_dist"):
            sample_vals = []
            for i in range(len(case["menus"])):
                r = build.LIST_FIRST + i
                ws = wb[build.LIST_SHEET]
                sample_vals.append([ws[f"{col}{r}"].value for col, *_ in build.LIST_COLS[1:]])
        results.append(dict(case=case, ok=ok, n=n, nhand=len(hand), fails=fails, sample=sample_vals))
        print(f"{case['id']} {'PASS' if ok else 'FAIL'} ({n} cells) {case['title']}")
        for f in fails:
            print("   ", f)
    shutil.rmtree(work, ignore_errors=True)

    lines = [
        "# P2 原価率計算シート — 計算テスト結果",
        "",
        f"- 実行日時: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"- 再計算に使ったソフト: {ver}（`soffice --headless --convert-to xlsx` で開いて保存し直し、再計算後の値を openpyxl で読んだ）",
        "- 期待値: (1) `tests/cases.py` に手で書いた主要セルの値、(2) `tests/model.py`（数式を使わず Python の Decimal で仕様どおり計算）による全計算セルの値",
        "- 判定: 数値は差が 1e-9 未満、文字は完全一致、空欄は空欄。さらに全シートにエラー値（#DIV/0!, #N/A など）が無いこと",
        "- 再実行: `python3 company/product/P2/tests/run_tests.py`",
        "",
        f"## 結果: {total_pass} 件合格 / {total_fail} 件不合格（全 {len(cases)} ケース）",
        "",
        "| ID | 内容 | 照合セル数（うち手計算） | 判定 |",
        "|---|---|---|---|",
    ]
    for r in results:
        c = r["case"]
        lines.append(f"| {c['id']} | {c['title']} | {r['n']}（{r['nhand']}） | {'合格' if r['ok'] else '**不合格**'} |")
    fails = [r for r in results if r["fails"]]
    if fails:
        lines += ["", "## 不合格の詳細", ""]
        for r in fails:
            lines.append(f"### {r['case']['id']}")
            lines += [f"- {f}" for f in r["fails"]]
    for r in results:
        if r["sample"]:
            lines += ["", "## 見本データの計算結果（実機確認で見比べる値）", "",
                      "`dist/genka-keisan.xlsx` の「一覧」シート。見本の架空の数値。", "",
                      "| メニュー名 | 販売価格（税込） | 合計原価（税抜） | 原価率 | 目標 | 目標との差 | 判定 | 推奨売価（税込） |",
                      "|---|---|---|---|---|---|---|---|"]
            for name, price, total, rate, target, diff, judge, rec in r["sample"]:
                lines.append(f"| {name} | {price:,} 円 | {total:,.2f} 円 | {rate:.2%} | {target:.1%} | "
                             f"{diff:+.2%} | {judge} | {rec:,.0f} 円 |")
    lines += ["", "## 確認していないこと", "",
              "- Microsoft Excel と Google スプレッドシートでの実機の計算結果（この環境には無い）。手順は `device-check.md`。"]
    with open(os.path.join(HERE, "result.md"), "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"\n合格 {total_pass} / 不合格 {total_fail}")
    return 0 if total_fail == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
