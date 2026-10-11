#!/usr/bin/env python3
"""アレルゲン一覧表の計算テスト。

1. tests/cases.py の各ケースの入力で、build.py と同じ数式のブックを作る
   （use_dist のケースは dist の同梱ファイルそのもの。edit_list があれば品目リストを書き換える）。
2. LibreOffice（soffice --headless --convert-to xlsx）で開いて保存し直し、数式を再計算させる。
3. 再計算後のセル値を読み、(a) ケースに手で書いた期待値、(b) model() で独立に計算した
   入力シート・お客様用シートの全50行・全計算セルの期待値と照合する。全シートにエラー値が無いことも確認する。
4. 結果を tests/result.md に書く（--no-write を付けると書かない。QA の再実行用）。不合格があれば終了コード 1。

使い方:  python3 company/product/P3/tests/run_tests.py
"""
from __future__ import annotations

import datetime
import os
import shutil
import subprocess
import sys
import tempfile

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
P3 = os.path.dirname(HERE)
sys.path.insert(0, P3)
sys.path.insert(0, HERE)

import build  # noqa: E402
from cases import CASES  # noqa: E402

ERROR_STRINGS = ("#DIV/0!", "#N/A", "#VALUE!", "#REF!", "#NAME?", "#NUM!", "Err:", "#NULL!")
C, N, Q = build.CONTAINS, build.NOTCONTAIN, build.UNKNOWN


def recalc(paths, outdir):
    cmd = ["soffice", "--headless", "--norestore", "--convert-to", "xlsx", "--outdir", outdir] + paths
    subprocess.run(cmd, check=True, stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=600)


def norm(v):
    if v is None:
        return ""
    if isinstance(v, datetime.datetime):
        return v.date()
    return v


def same(actual, expected):
    a, e = norm(actual), norm(expected)
    if isinstance(e, (int, float)) and not isinstance(e, bool):
        return isinstance(a, (int, float)) and abs(a - e) < 1e-9
    if isinstance(e, datetime.date):
        return a == e
    return str(a) == str(e)


def addr(key):
    """チェックキー → (シート名, セル番地)"""
    p = key.split(":")
    if p[0] == "cu" and p[1] == "title":
        return build.CUST_SHEET, "A1"
    if p[0] == "in" and p[1] == "sum_unk":
        return build.IN_SHEET, build.IN_SUM_UNK
    if p[0] == "in" and p[1] == "sum_todo":
        return build.IN_SHEET, build.IN_SUM_TODO
    if p[1] in ("head", "kubun"):
        i = int(p[2]) - 1
        if p[0] == "in":
            return build.IN_SHEET, f"{build.in_item_col(i)}{build.IN_HEAD if p[1] == 'head' else build.IN_HEAD_KUBUN}"
        return build.CUST_SHEET, f"{build.cu_item_col(i)}{build.CU_HEAD if p[1] == 'head' else build.CU_HEAD_KUBUN}"
    k = int(p[1]) - 1
    if p[0] == "in":
        col = {"status": build.IN_STATUS, "ncont": build.IN_NCONT, "nunk": build.IN_NUNK}[p[2]]
        return build.IN_SHEET, f"{col}{build.IN_FIRST + k}"
    r = build.CU_FIRST + k
    if p[2] == "item":
        names = [it[1] for it in build.ITEMS]
        return build.CUST_SHEET, f"{build.cu_item_col(names.index(p[3]))}{r}"
    col = {"name": "A", "date": build.CU_DATE, "note": build.CU_NOTE}[p[2]]
    return build.CUST_SHEET, f"{col}{r}"


# ---------------------------------------------------------------- 独立モデル（数式を使わずに仕様どおり計算）
def model_row(m):
    names = [it[1] for it in build.ITEMS]
    if m is None:
        m = dict(name=None, values={}, date=None, who=None, note=None)
    vs = [m["values"].get(n) for n in names]
    name = m.get("name")
    b = sum(1 for v in vs if v is None)
    q = sum(1 for v in vs if v == Q)
    c = sum(1 for v in vs if v == C)
    n = sum(1 for v in vs if v == N)
    empty = (not name) and b == 29
    if empty:
        st = ""
    elif not name:
        st = build.ST_NONAME
    elif c + n + q + b != 29:
        st = build.ST_BAD
    elif q + b > 0:
        st = build.ST_UNK
    elif m.get("date") is None:
        st = build.ST_NODATE
    elif m.get("who") is None:
        st = build.ST_NOWHO
    else:
        st = build.ST_OK
    cu = {}
    cu["name"] = name or ""
    cu["items"] = ["" if not name else (v if v in (C, N) else Q) for v in vs]
    cu["date"] = "" if (not name or m.get("date") is None) else m["date"]
    prefix = build.CU_UNK_PREFIX if st in (build.ST_UNK, build.ST_BAD) else ""
    cu["note"] = "" if not name else prefix + (m.get("note") or "")
    return dict(status=st, ncont="" if empty else c, nunk="" if empty else q + b, cu=cu)


def model_checks(menus, shop, items):
    out = []
    rows = [model_row(menus[k] if k < len(menus) else None) for k in range(build.N_MENUS)]
    for k, r in enumerate(rows, start=1):
        out += [(f"in:{k}:status", r["status"]), (f"in:{k}:ncont", r["ncont"]), (f"in:{k}:nunk", r["nunk"]),
                (f"cu:{k}:name", r["cu"]["name"]), (f"cu:{k}:date", r["cu"]["date"]),
                (f"cu:{k}:note", r["cu"]["note"])]
        for it, v in zip(build.ITEMS, r["cu"]["items"]):
            out.append((f"cu:{k}:item:{it[1]}", v))
    sts = [r["status"] for r in rows]
    out.append(("in:sum_unk", sts.count(build.ST_UNK)))
    out.append(("in:sum_todo", sum(1 for s in sts if s not in ("", build.ST_OK))))
    out.append(("cu:title", (shop + "　アレルゲン一覧表") if shop else "アレルゲン一覧表"))
    for i, (kubun, ja, *_r) in enumerate(items, start=1):
        out += [(f"in:head:{i}", ja), (f"cu:head:{i}", ja), (f"in:kubun:{i}", kubun), (f"cu:kubun:{i}", kubun)]
    return out


def scan_errors(wb):
    errs = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith(ERROR_STRINGS):
                    errs.append(f"{ws.title}!{c.coordinate}={c.value}")
    return errs


def prepare(case, path):
    """ケースの入力ブックを作り、model 用の (menus, shop, items) を返す"""
    if case.get("use_dist"):
        items = [list(it) for it in build.ITEMS]
        if case.get("edit_list"):
            wb = openpyxl.load_workbook(build.OUT)
            ws = wb[build.LIST_SHEET]
            for ref, v in case["edit_list"].items():
                ws[ref] = v
                i = int(ref[1:]) - build.LIST_FIRST
                items[i]["ABCDEF".index(ref[0]) - 1] = v
            wb.save(path)
        else:
            shutil.copy(build.OUT, path)
        return build.SAMPLE_MENUS, build.SAMPLE_SHOP, [tuple(it) for it in items]
    menus = [m if m is not None else dict(name=None, values={}, date=None, who=None, note=None)
             for m in case["menus"]]
    build.build_workbook(path, menus=menus, shop=case["shop"])
    return case["menus"], case["shop"], build.ITEMS


def main():
    work = tempfile.mkdtemp(prefix="p3test_")
    src, out = os.path.join(work, "src"), os.path.join(work, "out")
    os.makedirs(src)
    os.makedirs(out)
    paths, models = [], {}
    for case in CASES:
        p = os.path.join(src, f"{case['id']}.xlsx")
        models[case["id"]] = prepare(case, p)
        paths.append(p)
    recalc(paths, out)
    ver = subprocess.run(["soffice", "--version"], stdout=subprocess.PIPE, text=True).stdout.strip()

    results = []
    tp = tf = 0
    sample_rows = None
    for case in CASES:
        wb = openpyxl.load_workbook(os.path.join(out, f"{case['id']}.xlsx"), data_only=True)
        hand = case["checks"]
        auto = model_checks(*models[case["id"]])
        auto_map = dict(auto)
        fails = []
        for key, exp in hand:   # 手書きの期待値と model の食い違い（model の検算）
            if key in auto_map and not same(auto_map[key], exp):
                fails.append(f"model不一致 {key}: 手書き={exp!r} model={auto_map[key]!r}")
        n = 0
        for key, exp in hand + auto:
            sh, ref = addr(key)
            act = wb[sh][ref].value
            n += 1
            if not same(act, exp):
                fails.append(f"{key}（{sh}!{ref}）: 期待={exp!r} 実際={act!r}")
        fails += [f"エラー値: {e}" for e in scan_errors(wb)]
        ok = not fails
        tp += ok
        tf += not ok
        if case["id"] == "T15":
            ws_in, ws_cu = wb[build.IN_SHEET], wb[build.CUST_SHEET]
            sample_rows = []
            for k in range(3):
                r = build.IN_FIRST + k
                cr = build.CU_FIRST + k
                inc = [it[1] for i, it in enumerate(build.ITEMS)
                       if ws_cu[f"{build.cu_item_col(i)}{cr}"].value == C]
                unk = [it[1] for i, it in enumerate(build.ITEMS)
                       if ws_cu[f"{build.cu_item_col(i)}{cr}"].value == Q]
                sample_rows.append((ws_cu[f"A{cr}"].value, ws_in[f"{build.IN_STATUS}{r}"].value,
                                    "・".join(inc), "・".join(unk) or "なし"))
        results.append(dict(case=case, ok=ok, n=n, nhand=len(hand), fails=fails))
        print(f"{case['id']} {'PASS' if ok else 'FAIL'} ({n} cells) {case['title']}")
        for f in fails[:20]:
            print("   ", f)
    shutil.rmtree(work, ignore_errors=True)

    lines = [
        "# P3 アレルゲン一覧表 — 計算テスト結果",
        "",
        f"- 実行日時: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}",
        f"- 再計算に使ったソフト: {ver}（`soffice --headless --convert-to xlsx` で開いて保存し直し、再計算後の値を openpyxl で読んだ）",
        "- 期待値: (1) `tests/cases.py` に手で書いた主要セルの値、(2) `tests/run_tests.py` の model()"
        "（数式を使わず Python で仕様どおり計算）による、入力シートの状態・件数と、お客様用シートの全50行×（メニュー名・29品目・確認日・備考）、見出し、集計",
        "- 判定: 文字は完全一致、数値は一致、日付は日付として一致、空欄は空欄。さらに全シートにエラー値（#REF!, #NAME? など）が無いこと",
        "- 再実行: `python3 company/product/P3/tests/run_tests.py`",
        "",
        f"## 結果: {tp} 件合格 / {tf} 件不合格（全 {len(CASES)} ケース）",
        "",
        "| ID | 内容 | 照合セル数（うち手書き） | 判定 |",
        "|---|---|---|---|",
    ]
    for r in results:
        c = r["case"]
        lines.append(f"| {c['id']} | {c['title']} | {r['n']}（{r['nhand']}） | {'合格' if r['ok'] else '**不合格**'} |")
    bad = [r for r in results if r["fails"]]
    if bad:
        lines += ["", "## 不合格の詳細", ""]
        for r in bad:
            lines.append(f"### {r['case']['id']}")
            lines += [f"- {f}" for f in r["fails"][:50]]
    if sample_rows:
        lines += ["", "## 見本データの計算結果（実機確認で見比べる値）", "",
                  "`dist/allergen-ichiran.xlsx` の見本3品（見本の架空データ。実際のメニューには使えない）。", "",
                  "| メニュー名 | 入力シートの状態 | お客様用で ○ の品目 | お客様用で ？ の品目 |", "|---|---|---|---|"]
        for row in sample_rows:
            lines.append("| " + " | ".join(str(x) for x in row) + " |")
    lines += ["", "## 確認していないこと", "",
              "- Microsoft Excel と Google スプレッドシートでの実機の計算結果・プルダウン・縦書きの見出し・印刷（この環境には無い）。"
              "手順は `device-check.md`。",
              "- プルダウンで入力を制限できること（LibreOffice の再計算では入力規則は試せない）。入力規則が付いていることは `compat-check.md` で確認。"]
    if "--no-write" not in sys.argv[1:]:
        with open(os.path.join(HERE, "result.md"), "w", encoding="utf-8") as f:
            f.write("\n".join(lines) + "\n")
    else:
        print("（--no-write: result.md は書き換えていない）")
    print(f"\n合格 {tp} / 不合格 {tf}")
    return 0 if tf == 0 else 1


if __name__ == "__main__":
    sys.exit(main())
