#!/usr/bin/env python3
"""P3「アレルゲン一覧表」の xlsx を生成する。

使い方:
    python3 company/product/P3/build.py            # dist/allergen-ichiran.xlsx を作る（見本データ入り）

tests/ からは build_workbook(path, menus=..., shop=..., items=...) を呼び、
テスト用の入力を入れたブックを作る（数式は同じ）。

互換性の方針（Excel と Google スプレッドシート）:
- 関数は COUNTIF, COUNTBLANK, IF, AND, OR, ISBLANK だけを使う（許可リストは tests/compat_check.py）。
- 名前付き範囲（印刷タイトルを除く）・テーブル（構造化参照）・マクロ・外部リンク・シート保護は使わない。
- 入力規則のリストは「品目リスト」シートのセル範囲を参照する（直接文字列のリストは使わない）。
- 条件付き書式は「セルの値が ○ / ？ / 未確認あり に等しい」だけ。意味は必ず文字（記号）でも表示する。
"""
from __future__ import annotations

import datetime
import os
import sys

from openpyxl import Workbook
from openpyxl.cell.cell import MergedCell
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
DIST = os.path.join(HERE, "dist")
OUT = os.path.join(DIST, "allergen-ichiran.xlsx")

CAA_URL = "https://www.caa.go.jp/policies/policy/food_labeling/food_sanitation/allergy/index.html"

# ---- 品目（2026-04-01 改正後。research/2026-10-11-first-products.md §3） ----
# (区分, 日本語名, 英語名, 中国語簡体字名, 韓国語名)。訳は P1 の辞書と同じ（ピスタチオのみ追加）。
ITEMS = [
    ("義務", "えび", "Shrimp", "虾", "새우"),
    ("義務", "かに", "Crab", "蟹", "게"),
    ("義務", "くるみ", "Walnut", "核桃", "호두"),
    ("義務", "小麦", "Wheat", "小麦", "밀"),
    ("義務", "そば", "Buckwheat", "荞麦", "메밀"),
    ("義務", "卵", "Egg", "鸡蛋", "달걀"),
    ("義務", "乳", "Milk", "乳制品", "우유"),
    ("義務", "落花生（ピーナッツ）", "Peanut", "花生", "땅콩"),
    ("義務", "カシューナッツ", "Cashew nut", "腰果", "캐슈너트"),
    ("推奨", "アーモンド", "Almond", "杏仁", "아몬드"),
    ("推奨", "あわび", "Abalone", "鲍鱼", "전복"),
    ("推奨", "いか", "Squid", "鱿鱼", "오징어"),
    ("推奨", "いくら", "Salmon roe", "鲑鱼子", "연어알"),
    ("推奨", "オレンジ", "Orange", "橙子", "오렌지"),
    ("推奨", "キウイフルーツ", "Kiwi fruit", "猕猴桃", "키위"),
    ("推奨", "牛肉", "Beef", "牛肉", "쇠고기"),
    ("推奨", "ごま", "Sesame", "芝麻", "참깨"),
    ("推奨", "さけ", "Salmon", "三文鱼", "연어"),
    ("推奨", "さば", "Mackerel", "青花鱼", "고등어"),
    ("推奨", "ゼラチン", "Gelatin", "明胶", "젤라틴"),
    ("推奨", "大豆", "Soybean", "大豆", "대두"),
    ("推奨", "鶏肉", "Chicken", "鸡肉", "닭고기"),
    ("推奨", "バナナ", "Banana", "香蕉", "바나나"),
    ("推奨", "ピスタチオ", "Pistachio", "开心果", "피스타치오"),
    ("推奨", "豚肉", "Pork", "猪肉", "돼지고기"),
    ("推奨", "マカダミアナッツ", "Macadamia nut", "夏威夷果", "마카다미아"),
    ("推奨", "もも", "Peach", "桃子", "복숭아"),
    ("推奨", "やまいも", "Yam", "山药", "마"),
    ("推奨", "りんご", "Apple", "苹果", "사과"),
]
N_ITEMS = len(ITEMS)        # 29
N_DUTY = sum(1 for k, *_ in ITEMS if k == "義務")   # 9

CONTAINS, NOTCONTAIN, UNKNOWN = "○", "－", "？"   # U+25CB, U+FF0D, U+FF1F

# ---- レイアウト定数（tests からも参照する） ----
HOWTO_SHEET = "使い方"
LIST_SHEET = "品目リスト"
IN_SHEET = "入力"
CUST_SHEET = "お客様用"
IS = f"'{IN_SHEET}'"      # 数式で使うシート名（引用符付き）
LS = f"'{LIST_SHEET}'"

LIST_FIRST = 5                         # 品目リストのデータ開始行
LIST_LAST = LIST_FIRST + N_ITEMS - 1   # 33
CHOICE_COL = "I"                       # 入力の選択肢（○ － ？）
CHOICE_FIRST, CHOICE_LAST = 5, 7

N_MENUS = 50
IN_HEAD_KUBUN = 4                      # 入力：区分の行
IN_HEAD = 5                            # 入力：品目名の行
IN_FIRST = 6                           # 入力：データ開始行
IN_LAST = IN_FIRST + N_MENUS - 1       # 55
IN_SHOP = "B2"                         # 店名
IN_ITEM0 = 3                           # 入力：最初の品目の列番号（C）
IN_DATE = L(IN_ITEM0 + N_ITEMS)        # AF 最終確認日
IN_WHO = L(IN_ITEM0 + N_ITEMS + 1)     # AG 確認者
IN_NOTE = L(IN_ITEM0 + N_ITEMS + 2)    # AH 備考
IN_STATUS = L(IN_ITEM0 + N_ITEMS + 3)  # AI 状態
IN_NCONT = L(IN_ITEM0 + N_ITEMS + 4)   # AJ 含む品目の数
IN_NUNK = L(IN_ITEM0 + N_ITEMS + 5)    # AK 未確認の数（？＋空欄）
IN_SUM_UNK = "AI2"                     # 未確認ありの行数
IN_SUM_TODO = "AI3"                    # 確認済み以外（要確認）の行数

CU_HEAD_KUBUN = 4
CU_HEAD = 5
CU_FIRST = 6
CU_LAST = CU_FIRST + N_MENUS - 1
CU_ITEM0 = 2                           # お客様用：最初の品目の列番号（B）
CU_DATE = L(CU_ITEM0 + N_ITEMS)        # AE 確認日
CU_NOTE = L(CU_ITEM0 + N_ITEMS + 1)    # AF 備考
CU_LASTCOL = CU_NOTE


def in_item_col(i):
    return L(IN_ITEM0 + i)


def cu_item_col(i):
    return L(CU_ITEM0 + i)


# ---- 状態の文言 ----
ST_OK = "確認済み"
ST_UNK = "未確認あり"
ST_NODATE = "確認日なし"
ST_NOWHO = "確認者なし"
ST_BAD = "入力値エラー"
ST_NONAME = "メニュー名なし"
CU_UNK_PREFIX = "※未確認あり。"

# ---- 見本データ（架空） ----
SAMPLE_SHOP = "〇〇食堂（見本）"
SAMPLE_NOTE = "見本の架空データ。実際のメニューには使えない"


def _row(contains=(), unknown=(), blank=()):
    out = {}
    for k, name, *_ in ITEMS:
        if name in contains:
            out[name] = CONTAINS
        elif name in unknown:
            out[name] = UNKNOWN
        elif name in blank:
            out[name] = None
        else:
            out[name] = NOTCONTAIN
    return out


SAMPLE_MENUS = [
    dict(name="鶏の唐揚げ", values=_row(contains=("小麦", "卵", "大豆", "鶏肉")),
         date=datetime.date(2026, 10, 1), who="店長（見本）", note=SAMPLE_NOTE + "。揚げ油を共用"),
    dict(name="野菜サラダ（ごまドレッシング）", values=_row(contains=("卵", "ごま", "大豆", "りんご"), unknown=("乳",)),
         date=datetime.date(2026, 10, 1), who="店長（見本）", note=SAMPLE_NOTE),
    dict(name="ざるそば", values=_row(contains=("そば", "小麦", "大豆")),
         date=None, who=None, note=SAMPLE_NOTE),
]

# ---- 文言 ----
CUST_LEGEND = ("○＝含む　　－＝含まない　　？＝未確認（スタッフにおたずねください）　　"
               "区分　義務＝特定原材料（9品目）　推奨＝特定原材料に準ずるもの（20品目）")
CUST_DISCLAIMER = (
    "この表は、当店がお客様に自主的にお知らせするものです（飲食店のアレルギー表示は法律上の義務ではありません）。"
    "内容は当店が原材料を確認して作成していますが、仕入れ先や原材料の変更により変わることがあります。確認した日は「確認日」欄をご覧ください。"
    "調理器具・揚げ油などを共用しているため、「－」の品目が混ざる可能性があります。"
    "重いアレルギーのあるお客様は、ご注文の前に必ずスタッフにご相談ください。"
    "対象は消費者庁の資料による特定原材料等29品目（2026年4月時点）です。")

HOWTO_LINES = [
    ("アレルゲン一覧表（特定原材料等29品目）", "title"),
    ("お店のメニューごとに、アレルギーの原因となる食品（消費者庁の資料による特定原材料等の29品目）が入っているかを一覧にして、お客様に見せるための表です。", None),
    ("", None),
    ("■ 必ずお読みください", "head"),
    ("1. 飲食店のアレルギー表示は、法律上の義務ではありません。この表は、お店が自主的にお客様へ情報を提供するためのものです。", None),
    ("2. 表の内容（○・－・？）の最終確認は、お店（この表をお使いの方）が行ってください。制作者は、個々のメニューの原材料を確認していません。", None),
    ("3. 仕入れ先・原材料・レシピを変えたら、そのたびに表を更新し、「最終確認日」を書き換えてください。同じ名前の食品でも、メーカーによって原材料が違うことがあります。", None),
    ("4. 調理器具・揚げ油・まな板などの共用により、「－」の品目が混ざる可能性があります。お客様用の表にもその旨を書いています。具体的な共用（例：揚げ油を共用）は「備考」に書いてください。", None),
    ("5. 重いアレルギーのあるお客様には、表だけで判断せず、スタッフに相談していただくよう伝えてください。", None),
    ("6. 29品目以外の食品にアレルギーのある方もいます。表にない食品は、お客様ごとに確認してください。", None),
    ("7. 品目は2026年4月時点の消費者庁の資料に基づいています（特定原材料9品目＋特定原材料に準ずるもの20品目）。最新の品目は消費者庁のページで確認してください。", None),
    ("　消費者庁「食物アレルギー表示に関する情報」 " + CAA_URL, None),
    ("", None),
    ("■ シートの説明", "head"),
    ("・品目リスト … 29品目の区分（義務／推奨）・日本語名・英語名・中国語（簡体字）名・韓国語名。ここを直すと、入力とお客様用の見出しに自動で反映されます。", None),
    ("・入力 … メニューごとに29品目を ○／－／？ で入力します（最大50品）。お店の中で使うシートです。", None),
    ("・お客様用 … 入力の内容から自動で作られる、お客様に見せる表です（A4横で印刷）。このシートには入力しないでください。", None),
    ("", None),
    ("■ 使い方", "head"),
    ("1. 「入力」シートの B2 に店名を入れます（お客様用の表題になります）。", None),
    ("2. 見本の3行（鶏の唐揚げ・野菜サラダ・ざるそば）の内容を消します。行の削除はしないでください（セルを選んで Delete キーで内容だけ消す）。行を削除すると、お客様用の表がずれたりエラーになったりします。", None),
    ("3. メニュー名を上から詰めて入力します（途中に空いた行があると、お客様用にも空いた行が出ます）。", None),
    ("4. 29品目それぞれに、プルダウンで ○（含む）／－（含まない）／？（未確認）を選びます。わからないものは「？」にして、あとで原材料の表示や仕入れ先の資料で確認します。", None),
    ("5. 「最終確認日」（例：2026/10/11）と「確認者」を入れます。「備考」には調理器具・揚げ油の共用などを書きます。", None),
    ("6. 右端の「状態」がすべて「確認済み」になったら、「お客様用」シートを印刷します。", None),
    ("", None),
    ("■ 「状態」の意味（入力シートの右端）", "head"),
    ("・確認済み … 29品目すべてが ○ か －、確認日と確認者も入っている", None),
    ("・未確認あり … 「？」または空欄が残っている。お客様用では、その品目が「？」になり、備考に「※未確認あり。」と出ます", None),
    ("・確認日なし／確認者なし … 最終確認日・確認者が空欄", None),
    ("・入力値エラー … ○ － ？ 以外の文字が入っている（貼り付けなど）。プルダウンで選び直してください。お客様用では「？」になります", None),
    ("・メニュー名なし … 品目は入っているのに、メニュー名が空欄（お客様用には出ません）", None),
    ("", None),
    ("■ 印刷（お客様用シート）", "head"),
    ("・Excel：お客様用シートを開き、ファイル → 印刷。用紙 A4・横・拡大縮小なし（100%）。見出しと注意書きは各ページに印刷されます。メニューが少ないときは、印刷するページを指定してください（例：1ページ目だけ）。", None),
    ("・Google スプレッドシート：ファイル → 印刷 →「現在のシート」、用紙サイズ A4、ページの向き 横、スケール「標準（100%）」、「ヘッダーとフッター」の「固定行を繰り返す」をオン。", None),
    ("・文字は印刷時に10ポイント以上になるように作っています。縮小して印刷すると小さくなります。", None),
    ("", None),
    ("■ 品目を直すとき", "head"),
    ("・品目の名前・区分・訳は「品目リスト」シートだけを直してください。入力とお客様用の見出しに自動で反映されます。", None),
    ("・品目の数（29）が変わる改正があった場合は、列の追加が必要なため、このファイルでは自動で対応できません。", None),
    ("・英語・中国語（簡体字）・韓国語の名前は参考訳です。掲示に使う前に、ご自身または詳しい方が確認してください。", None),
    ("", None),
    ("■ Google スプレッドシートで使うとき", "head"),
    ("・Google ドライブにアップロードし、「Google スプレッドシートで開く」→ ファイル →「Google スプレッドシートとして保存」。", None),
    ("・スマートフォンでもプルダウンから選べます（プルダウンの作成・変更はパソコンが必要です）。", None),
    ("", None),
    ("■ その他", "head"),
    ("・色の付いたセル（区分・品目名・状態・お客様用シート）は自動計算です。消したり上書きしたりしないでください。シートの保護はしていません。", None),
    ("・見本の店名「〇〇食堂（見本）」と3品のメニューは架空のものです。見本の ○／－ は「見本の架空データ。実際のメニューには使えない」ものです。", None),
]

# ---- スタイル ----
FONT = "Meiryo"
F_TITLE = Font(name=FONT, size=14, bold=True)
F_HEAD = Font(name=FONT, size=12, bold=True)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_NORMAL = Font(name=FONT, size=10)
F_WARN = Font(name=FONT, size=10, bold=True, color="C00000")
F_BIG = Font(name=FONT, size=11, bold=True)
FILL_HEAD = PatternFill("solid", fgColor="DDEBF7")
FILL_DUTY = PatternFill("solid", fgColor="FCE4D6")
FILL_CALC = PatternFill("solid", fgColor="F2F2F2")
FILL_INPUT = PatternFill("solid", fgColor="FFFFFF")
FILL_Q = PatternFill("solid", fgColor="FFF2CC")
FILL_C = PatternFill("solid", fgColor="F8CBAD")
FILL_WARN = PatternFill("solid", fgColor="FFD966")
THIN = Side(style="thin", color="808080")
THICK = Side(style="medium", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP = Alignment(vertical="center", wrap_text=True)
TOPWRAP = Alignment(vertical="top", wrap_text=True)
VERTICAL = Alignment(horizontal="center", vertical="bottom", text_rotation=255, wrap_text=False)


def _cell(ws, ref, value=None, font=F_NORMAL, fill=None, fmt=None, border=True, align=None):
    c = ws[ref]
    if not isinstance(c, MergedCell):
        c.value = value
    c.font = font
    if fill is not None:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if border:
        c.border = BORDER
    if align is not None:
        c.alignment = align
    return c


# ---------------------------------------------------------------- 使い方
def _howto(ws):
    ws.column_dimensions["A"].width = 110
    for i, (text, kind) in enumerate(HOWTO_LINES, start=1):
        font = {"title": F_TITLE, "head": F_HEAD}.get(kind, F_NORMAL)
        if text.startswith(("1. 飲食店", "2. 表の内容", "5. 重い")):
            font = F_BOLD
        _cell(ws, f"A{i}", text, font, border=False, align=TOPWRAP)
    ws.sheet_view.showGridLines = False
    ws.page_setup.orientation = "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ---------------------------------------------------------------- 品目リスト
def _itemlist(ws, items):
    _cell(ws, "A1", "品目リスト（ここを直すと、入力・お客様用の見出しに反映されます）", F_TITLE, border=False)
    _cell(ws, "A2", "品目は2026年4月時点の消費者庁の資料に基づく。最新の品目は消費者庁のページで確認してください。 "
          + CAA_URL, F_WARN, border=False)
    _cell(ws, "A3", "英語・中国語（簡体字）・韓国語の名前は参考訳です（ネイティブの確認はしていません）。"
          "区分：義務＝特定原材料（容器包装された加工食品で表示が義務の品目）、推奨＝特定原材料に準ずるもの（表示が推奨の品目）。",
          F_NORMAL, border=False)
    heads = [("A", "No", 5), ("B", "区分", 7), ("C", "日本語名", 22), ("D", "英語名（参考訳）", 18),
             ("E", "中国語簡体字名（参考訳）", 18), ("F", "韓国語名（参考訳）", 16)]
    for col, text, w in heads:
        _cell(ws, f"{col}4", text, F_BOLD, FILL_HEAD, align=CENTER)
        ws.column_dimensions[col].width = w
    for i, (kubun, ja, en, zh, ko) in enumerate(items):
        r = LIST_FIRST + i
        fill = FILL_DUTY if kubun == "義務" else None
        _cell(ws, f"A{r}", i + 1, align=CENTER, fill=fill)
        _cell(ws, f"B{r}", kubun, align=CENTER, fill=fill)
        _cell(ws, f"C{r}", ja, fill=fill)
        _cell(ws, f"D{r}", en, fill=fill)
        _cell(ws, f"E{r}", zh, fill=fill)
        _cell(ws, f"F{r}", ko, fill=fill)
    ws.column_dimensions["H"].width = 2
    ws.column_dimensions[CHOICE_COL].width = 10
    ws.column_dimensions["J"].width = 26
    _cell(ws, f"{CHOICE_COL}4", "入力の選択肢", F_BOLD, FILL_HEAD, align=CENTER)
    _cell(ws, "J4", "意味（この3つは変えないでください）", F_BOLD, FILL_HEAD, align=CENTER)
    for r, (sym, mean) in zip(range(CHOICE_FIRST, CHOICE_LAST + 1),
                              [(CONTAINS, "含む"), (NOTCONTAIN, "含まない"), (UNKNOWN, "未確認")]):
        _cell(ws, f"{CHOICE_COL}{r}", sym, F_BIG, align=CENTER)
        _cell(ws, f"J{r}", mean)
    ws.freeze_panes = "A5"


# ---------------------------------------------------------------- 入力
def _status_formula(r):
    rng = f"{in_item_col(0)}{r}:{in_item_col(N_ITEMS - 1)}{r}"
    b = f"COUNTBLANK({rng})"
    q = f'COUNTIF({rng},"{UNKNOWN}")'
    c = f'COUNTIF({rng},"{CONTAINS}")'
    n = f'COUNTIF({rng},"{NOTCONTAIN}")'
    return (f'=IF(AND($B{r}="",{b}={N_ITEMS}),"",IF($B{r}="","{ST_NONAME}",'
            f'IF({c}+{n}+{q}+{b}<>{N_ITEMS},"{ST_BAD}",'
            f'IF({q}+{b}>0,"{ST_UNK}",'
            f'IF(ISBLANK({IN_DATE}{r}),"{ST_NODATE}",'
            f'IF(ISBLANK({IN_WHO}{r}),"{ST_NOWHO}","{ST_OK}"))))))')


def _input(ws, menus, shop):
    last_item = in_item_col(N_ITEMS - 1)
    _cell(ws, "A1", "入力（お店の中で使うシート）", F_TITLE, border=False)
    _cell(ws, "A2", "店名", F_BOLD, FILL_HEAD, align=CENTER)
    _cell(ws, IN_SHOP, shop, F_BIG, FILL_INPUT)
    _cell(ws, "C2", "← 店名（お客様用の表題になります）", border=False)
    _cell(ws, "A3", "見本の3行は「" + SAMPLE_NOTE + "」。内容を消して（行は削除しない）お使いください。"
          "入力は ○（含む）／－（含まない）／？（未確認）。", F_WARN, border=False)
    ws.merge_cells(f"{IN_DATE}2:{IN_NOTE}2")
    _cell(ws, f"{IN_DATE}2", "未確認ありの行数 →", F_BOLD, border=False, align=Alignment(horizontal="right"))
    _cell(ws, IN_SUM_UNK, f'=COUNTIF({IN_STATUS}{IN_FIRST}:{IN_STATUS}{IN_LAST},"{ST_UNK}")', F_BIG, FILL_CALC, align=CENTER)
    ws.merge_cells(f"{IN_DATE}3:{IN_NOTE}3")
    _cell(ws, f"{IN_DATE}3", "確認済み以外の行数 →", F_BOLD, border=False, align=Alignment(horizontal="right"))
    sr = f"{IN_STATUS}{IN_FIRST}:{IN_STATUS}{IN_LAST}"
    _cell(ws, IN_SUM_TODO,
          f'=COUNTIF({sr},"{ST_UNK}")+COUNTIF({sr},"{ST_NODATE}")+COUNTIF({sr},"{ST_NOWHO}")'
          f'+COUNTIF({sr},"{ST_BAD}")+COUNTIF({sr},"{ST_NONAME}")', F_BIG, FILL_CALC, align=CENTER)

    # 見出し
    for ref, text in [("A", "No"), ("B", "メニュー名")]:
        ws.merge_cells(f"{ref}{IN_HEAD_KUBUN}:{ref}{IN_HEAD}")
        _cell(ws, f"{ref}{IN_HEAD_KUBUN}", text, F_BOLD, FILL_HEAD, align=CENTER)
        _cell(ws, f"{ref}{IN_HEAD}", None, F_BOLD, FILL_HEAD)
    for i in range(N_ITEMS):
        col = in_item_col(i)
        lr = LIST_FIRST + i
        _cell(ws, f"{col}{IN_HEAD_KUBUN}", f"={LS}!$B${lr}", F_BOLD, FILL_CALC, align=CENTER)
        _cell(ws, f"{col}{IN_HEAD}", f"={LS}!$C${lr}", F_BOLD, FILL_CALC, align=VERTICAL)
        ws.column_dimensions[col].width = 4.2
    tail = [(IN_DATE, "最終確認日", 12), (IN_WHO, "確認者", 10), (IN_NOTE, "備考（調理器具・揚げ油の共用など）", 30),
            (IN_STATUS, "状態［自動］", 13), (IN_NCONT, "含む品目の数［自動］", 8), (IN_NUNK, "未確認の数［自動］", 8)]
    for col, text, w in tail:
        ws.merge_cells(f"{col}{IN_HEAD_KUBUN}:{col}{IN_HEAD}")
        fill = FILL_CALC if "自動" in text else FILL_HEAD
        _cell(ws, f"{col}{IN_HEAD_KUBUN}", text, F_BOLD, fill, align=CENTER)
        _cell(ws, f"{col}{IN_HEAD}", None, F_BOLD, fill)
        ws.column_dimensions[col].width = w
    ws.column_dimensions["A"].width = 5
    ws.column_dimensions["B"].width = 26
    ws.row_dimensions[IN_HEAD].height = 150
    ws.row_dimensions[3].height = 20

    for k in range(N_MENUS):
        r = IN_FIRST + k
        m = menus[k] if k < len(menus) else None
        rng = f"{in_item_col(0)}{r}:{last_item}{r}"
        _cell(ws, f"A{r}", k + 1, align=CENTER, fill=FILL_CALC)
        _cell(ws, f"B{r}", m.get("name") if m else None, align=WRAP)
        for i, (_, name, *_r) in enumerate(ITEMS):
            v = m["values"].get(name) if m else None
            _cell(ws, f"{in_item_col(i)}{r}", v, F_BIG, align=CENTER)
        _cell(ws, f"{IN_DATE}{r}", m.get("date") if m else None, fmt="yyyy/mm/dd", align=CENTER)
        _cell(ws, f"{IN_WHO}{r}", m.get("who") if m else None, align=WRAP)
        _cell(ws, f"{IN_NOTE}{r}", m.get("note") if m else None, align=WRAP)
        _cell(ws, f"{IN_STATUS}{r}", _status_formula(r), F_BOLD, FILL_CALC, align=CENTER)
        empty = f'AND($B{r}="",COUNTBLANK({rng})={N_ITEMS})'
        _cell(ws, f"{IN_NCONT}{r}", f'=IF({empty},"",COUNTIF({rng},"{CONTAINS}"))', fill=FILL_CALC, align=CENTER)
        _cell(ws, f"{IN_NUNK}{r}", f'=IF({empty},"",COUNTIF({rng},"{UNKNOWN}")+COUNTBLANK({rng}))',
              fill=FILL_CALC, align=CENTER)
        ws.row_dimensions[r].height = 30
        # 義務と推奨の境目に太線
        c = ws[f"{in_item_col(N_DUTY - 1)}{r}"]
        c.border = Border(left=THIN, right=THICK, top=THIN, bottom=THIN)

    # プルダウン（品目リストの選択肢を参照）
    dv = DataValidation(type="list", formula1=f"{LS}!${CHOICE_COL}${CHOICE_FIRST}:${CHOICE_COL}${CHOICE_LAST}",
                        allow_blank=True, showDropDown=False, showErrorMessage=True, errorStyle="stop",
                        errorTitle="入力できる値", error="○（含む）、－（含まない）、？（未確認）から選んでください。",
                        showInputMessage=False)
    dv.add(f"{in_item_col(0)}{IN_FIRST}:{last_item}{IN_LAST}")
    ws.add_data_validation(dv)
    area = f"{in_item_col(0)}{IN_FIRST}:{last_item}{IN_LAST}"
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[f'"{UNKNOWN}"'], fill=FILL_Q))
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[f'"{CONTAINS}"'], fill=FILL_C))
    ws.conditional_formatting.add(f"{IN_STATUS}{IN_FIRST}:{IN_STATUS}{IN_LAST}",
                                  CellIsRule(operator="equal", formula=[f'"{ST_UNK}"'], fill=FILL_WARN))
    ws.freeze_panes = f"C{IN_FIRST}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# ---------------------------------------------------------------- お客様用
def _customer(ws):
    last = CU_LASTCOL
    ws.merge_cells(f"A1:{last}1")
    _cell(ws, "A1", f'=IF(ISBLANK({IS}!${IN_SHOP[0]}${IN_SHOP[1:]}),"アレルゲン一覧表",'
                    f'{IS}!${IN_SHOP[0]}${IN_SHOP[1:]}&"　アレルゲン一覧表")',
          F_TITLE, border=False, align=Alignment(horizontal="left", vertical="center"))
    ws.merge_cells(f"A2:{last}2")
    _cell(ws, "A2", CUST_LEGEND, F_BOLD, border=False, align=Alignment(vertical="center"))
    ws.merge_cells(f"A3:{last}3")
    _cell(ws, "A3", CUST_DISCLAIMER, F_NORMAL, border=False, align=TOPWRAP)
    ws.row_dimensions[1].height = 24
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 54

    _cell(ws, f"A{CU_HEAD_KUBUN}", "区分", F_BOLD, FILL_HEAD, align=CENTER)
    _cell(ws, f"A{CU_HEAD}", "メニュー", F_BOLD, FILL_HEAD, align=CENTER)
    for i in range(N_ITEMS):
        col = cu_item_col(i)
        lr = LIST_FIRST + i
        _cell(ws, f"{col}{CU_HEAD_KUBUN}", f"={LS}!$B${lr}", F_NORMAL, FILL_HEAD,
              align=Alignment(horizontal="center", vertical="center", text_rotation=255))
        _cell(ws, f"{col}{CU_HEAD}", f"={LS}!$C${lr}", F_BOLD, FILL_HEAD, align=VERTICAL)
    for col, text in [(CU_DATE, "確認日"), (CU_NOTE, "備考")]:
        ws.merge_cells(f"{col}{CU_HEAD_KUBUN}:{col}{CU_HEAD}")
        _cell(ws, f"{col}{CU_HEAD_KUBUN}", text, F_BOLD, FILL_HEAD, align=CENTER)
        _cell(ws, f"{col}{CU_HEAD}", None, F_BOLD, FILL_HEAD)
    ws.row_dimensions[CU_HEAD_KUBUN].height = 30
    ws.row_dimensions[CU_HEAD].height = 140

    for k in range(N_MENUS):
        r = CU_FIRST + k
        ir = IN_FIRST + k
        _cell(ws, f"A{r}", f'=IF(ISBLANK({IS}!$B{ir}),"",{IS}!$B{ir})', F_NORMAL, align=WRAP)
        for i in range(N_ITEMS):
            src = f"{IS}!{in_item_col(i)}{ir}"
            _cell(ws, f"{cu_item_col(i)}{r}",
                  f'=IF($A{r}="","",IF({src}="{CONTAINS}","{CONTAINS}",IF({src}="{NOTCONTAIN}","{NOTCONTAIN}","{UNKNOWN}")))',
                  F_BIG, align=CENTER)
        _cell(ws, f"{CU_DATE}{r}", f'=IF(OR($A{r}="",ISBLANK({IS}!{IN_DATE}{ir})),"",{IS}!{IN_DATE}{ir})',
              F_NORMAL, fmt="yyyy/m/d", align=CENTER)
        _cell(ws, f"{CU_NOTE}{r}",
              f'=IF($A{r}="","",IF(OR({IS}!{IN_STATUS}{ir}="{ST_UNK}",{IS}!{IN_STATUS}{ir}="{ST_BAD}"),'
              f'"{CU_UNK_PREFIX}","")&{IS}!{IN_NOTE}{ir})', F_NORMAL, align=WRAP)
        ws.row_dimensions[r].height = 28
    # 義務と推奨の境目に太線（見出しから最終行まで）
    bc = cu_item_col(N_DUTY - 1)
    for r in range(CU_HEAD_KUBUN, CU_LAST + 1):
        c = ws[f"{bc}{r}"]
        c.border = Border(left=THIN, right=THICK, top=THIN, bottom=THIN)
    area = f"{cu_item_col(0)}{CU_FIRST}:{cu_item_col(N_ITEMS - 1)}{CU_LAST}"
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[f'"{UNKNOWN}"'], fill=FILL_Q))
    ws.conditional_formatting.add(area, CellIsRule(operator="equal", formula=[f'"{CONTAINS}"'], fill=FILL_C))

    ws.column_dimensions["A"].width = 17
    for i in range(N_ITEMS):
        ws.column_dimensions[cu_item_col(i)].width = 3.0
    ws.column_dimensions[CU_DATE].width = 9.2
    ws.column_dimensions[CU_NOTE].width = 30

    ws.freeze_panes = f"A{CU_FIRST}"
    ws.print_title_rows = f"1:{CU_HEAD}"
    ws.print_area = f"A1:{last}{CU_LAST}"
    ps = ws.page_setup
    ps.orientation = "landscape"
    ps.paperSize = ws.PAPERSIZE_A4
    ps.scale = 100
    ws.sheet_properties.pageSetUpPr.fitToPage = False
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.25
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.45
    ws.page_margins.header = 0.15
    ws.page_margins.footer = 0.2
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.center.size = 10


# ---------------------------------------------------------------- 全体
def build_workbook(path, menus=None, shop=None, items=None):
    menus = SAMPLE_MENUS if menus is None else menus
    shop = SAMPLE_SHOP if shop is None else shop
    items = ITEMS if items is None else items
    wb = Workbook()
    ws = wb.active
    ws.title = HOWTO_SHEET
    _howto(ws)
    _itemlist(wb.create_sheet(LIST_SHEET), items)
    _input(wb.create_sheet(IN_SHEET), menus, shop)
    _customer(wb.create_sheet(CUST_SHEET))
    wb.calculation = CalcProperties(fullCalcOnLoad=True)
    wb.security = None
    wb.properties.creator = "MenuBridge"
    wb.properties.title = "アレルゲン一覧表"
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else OUT
    print(build_workbook(out))
